"""PDF / Word export for the Report Agent (Module 7 addendum).

The single source of truth for export is ``FinalReport.markdown`` - the fully
assembled document string already produced by ``run_report``/
``assemble_markdown`` (header + all ten sections with LLM intros woven in +
deterministic tables). We do NOT re-derive content from ``FinalReport.sections``
- those don't carry the LLM ``section_intros`` prose, so re-deriving from them
would produce an incomplete export.

This is a small, deliberately limited Markdown-to-document converter - not a
general-purpose Markdown parser, just enough to handle the exact subset this
codebase's own renderers ever produce (``_fmt``, ``_md_table``, ``_bullets``,
``_stub``, ``_blockquote`` in ``app/agents/report.py``, plus the ``# ``/``## ``
headers ``run_report``/``assemble_markdown`` emit).

``parse_markdown_blocks`` turns the markdown string into a shared intermediate
representation (a list of small tagged-union dicts). Two renderers then walk
that SAME block list - one with ``reportlab`` flowables into a PDF, the other
with ``python-docx`` into a ``.docx`` - so the two output formats can't drift
out of sync with each other.
"""

from __future__ import annotations

import io
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.agents.report import FinalReport

Block = dict[str, Any]


# ---------------------------------------------------------------------------
# Markdown -> block list
# ---------------------------------------------------------------------------
def _is_pipe_row(stripped: str) -> bool:
    return stripped.startswith("|") and stripped.endswith("|") and len(stripped) >= 2


def _is_separator_row(stripped: str) -> bool:
    cells = [c.strip() for c in stripped.strip("|").split("|")]
    return bool(cells) and all(c and set(c) <= set(":-") for c in cells)


def _is_italic_line(stripped: str) -> bool:
    """``_..._``: starts and ends with a single underscore, nothing else."""
    if len(stripped) < 2:
        return False
    if stripped[0] != "_" or stripped[-1] != "_":
        return False
    if len(stripped) >= 3 and (stripped[1] == "_" or stripped[-2] == "_"):
        return False
    return True


def _is_bold_subheading(stripped: str) -> bool:
    """``**text**`` alone on its own line (nothing else on the line)."""
    if len(stripped) < 5 or not (stripped.startswith("**") and stripped.endswith("**")):
        return False
    inner = stripped[2:-2]
    return bool(inner) and "**" not in inner


def parse_markdown_blocks(markdown: str) -> list[Block]:
    """Classify ``markdown`` line-by-line into a shared block list.

    Per-line classification, in order (see the Task 6 brief for the full
    rationale of each rule):

    1. ``## `` prefix -> section heading.
    2. ``# `` prefix -> document title.
    3. ``**text**`` alone on its own line -> bold sub-heading.
    4. ``| ... |`` -> GitHub-style pipe table row (accumulated).
    5. ``> `` prefix -> blockquote paragraph (accumulated).
    6. ``- `` prefix -> bullet list item (accumulated).
    7. Exactly ``_..._`` -> italic paragraph.
    8. ``_Generated <timestamp> (UTC)_`` -> matches rule 7.
    9. Blank line -> spacer.
    10. Anything else -> plain paragraph.
    """
    lines = markdown.split("\n")
    blocks: list[Block] = []
    i = 0
    n = len(lines)

    while i < n:
        raw = lines[i]
        stripped = raw.strip()

        if stripped.startswith("## "):
            blocks.append({"type": "heading", "text": stripped[3:].strip()})
            i += 1
            continue

        if stripped.startswith("# "):
            blocks.append({"type": "title", "text": stripped[2:].strip()})
            i += 1
            continue

        if _is_bold_subheading(stripped):
            blocks.append({"type": "subheading", "text": stripped[2:-2].strip()})
            i += 1
            continue

        if _is_pipe_row(stripped):
            table_lines: list[str] = []
            while i < n and _is_pipe_row(lines[i].strip()):
                table_lines.append(lines[i].strip())
                i += 1
            header = [c.strip() for c in table_lines[0].strip("|").split("|")]
            body_lines = table_lines[1:]
            if body_lines and _is_separator_row(body_lines[0]):
                body_lines = body_lines[1:]
            rows = [
                [c.strip() for c in line.strip("|").split("|")]
                for line in body_lines
            ]
            blocks.append({"type": "table", "headers": header, "rows": rows})
            continue

        if stripped.startswith("> "):
            quote_lines: list[str] = []
            while i < n and lines[i].strip().startswith("> "):
                quote_lines.append(lines[i].strip()[2:])
                i += 1
            blocks.append({"type": "blockquote", "text": "\n".join(quote_lines)})
            continue

        if stripped.startswith("- "):
            items: list[str] = []
            while i < n and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:])
                i += 1
            blocks.append({"type": "bullet", "items": items})
            continue

        if _is_italic_line(stripped):
            blocks.append({"type": "italic", "text": stripped[1:-1]})
            i += 1
            continue

        if stripped == "":
            blocks.append({"type": "spacer"})
            i += 1
            continue

        blocks.append({"type": "paragraph", "text": stripped})
        i += 1

    return blocks


# ---------------------------------------------------------------------------
# PDF renderer (reportlab)
# ---------------------------------------------------------------------------
def render_pdf(final_report: FinalReport) -> bytes:
    blocks = parse_markdown_blocks(final_report.markdown)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=LETTER,
        leftMargin=0.75 * inch, rightMargin=0.75 * inch,
        topMargin=0.75 * inch, bottomMargin=0.75 * inch,
    )
    styles = getSampleStyleSheet()
    subheading_style = ParagraphStyle(
        "SubHeading", parent=styles["Heading3"], spaceAfter=4,
    )
    quote_style = ParagraphStyle(
        "Quote", parent=styles["Italic"], leftIndent=18, spaceAfter=6,
    )

    flowables: list[Any] = []
    for block in blocks:
        btype = block["type"]
        if btype == "title":
            flowables.append(Paragraph(escape(block["text"]), styles["Title"]))
        elif btype == "heading":
            flowables.append(Paragraph(escape(block["text"]), styles["Heading2"]))
        elif btype == "subheading":
            flowables.append(Paragraph(escape(block["text"]), subheading_style))
        elif btype == "table":
            headers = block["headers"]
            rows = block["rows"]
            data = [headers, *rows] if rows else [headers]
            table = Table(data, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            flowables.append(table)
            flowables.append(Spacer(1, 6))
        elif btype == "blockquote":
            flowables.append(Paragraph(escape(block["text"]).replace("\n", "<br/>"), quote_style))
        elif btype == "bullet":
            items = [
                ListItem(Paragraph(escape(item), styles["Normal"]))
                for item in block["items"]
            ]
            flowables.append(ListFlowable(items, bulletType="bullet"))
        elif btype == "italic":
            flowables.append(Paragraph(f"<i>{escape(block['text'])}</i>", styles["Normal"]))
        elif btype == "spacer":
            flowables.append(Spacer(1, 6))
        else:  # paragraph
            flowables.append(Paragraph(escape(block["text"]), styles["Normal"]))

    doc.build(flowables)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Word renderer (python-docx)
# ---------------------------------------------------------------------------
def render_docx(final_report: FinalReport) -> bytes:
    blocks = parse_markdown_blocks(final_report.markdown)
    document = Document()

    for block in blocks:
        btype = block["type"]
        if btype == "title":
            document.add_heading(block["text"], level=0)
        elif btype == "heading":
            document.add_heading(block["text"], level=1)
        elif btype == "subheading":
            p = document.add_paragraph()
            p.add_run(block["text"]).bold = True
        elif btype == "table":
            headers = block["headers"]
            rows = block["rows"]
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [
                s.name for s in document.styles
            ] else "Table Grid"
            hdr_cells = table.rows[0].cells
            for idx, h in enumerate(headers):
                hdr_cells[idx].paragraphs[0].add_run(h).bold = True
            for row in rows:
                cells = table.add_row().cells
                for idx, value in enumerate(row):
                    if idx < len(cells):
                        cells[idx].text = value
        elif btype == "blockquote":
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = None
            run = p.add_run(block["text"])
            run.italic = True
        elif btype == "bullet":
            for item in block["items"]:
                document.add_paragraph(item, style="List Bullet")
        elif btype == "italic":
            p = document.add_paragraph()
            p.add_run(block["text"]).italic = True
        elif btype == "spacer":
            continue
        else:  # paragraph
            document.add_paragraph(block["text"])

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
