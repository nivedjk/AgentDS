"""Tests for PDF/Word report export (Module 7 addendum)."""

import io

import pytest
from docx import Document

from app.agents.report import FinalReport, run_report
from app.agents.report_export import parse_markdown_blocks, render_docx, render_pdf
from app.storage.dataset_store import save_artifact
from tests.test_report import (
    FAKE_CLEANING,
    FAKE_TRAINING,
    FAKE_UNDERSTANDING,
    PROSE,
    FakeAnthropicClient,
    _response,
    _tool_use_block,
)


@pytest.fixture(autouse=True)
def _data_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENTDS_DATA_DIR", str(tmp_path / "uploads"))


# ---------------------------------------------------------------------------
# parse_markdown_blocks unit tests - one per line-classification rule
# ---------------------------------------------------------------------------
def test_parse_document_title():
    blocks = parse_markdown_blocks("# AgentDS Final Report — ds-1")
    assert blocks == [{"type": "title", "text": "AgentDS Final Report — ds-1"}]


def test_parse_section_heading():
    blocks = parse_markdown_blocks("## Cleaning")
    assert blocks == [{"type": "heading", "text": "Cleaning"}]


def test_parse_bold_subheading():
    blocks = parse_markdown_blocks("**Key findings**")
    assert blocks == [{"type": "subheading", "text": "Key findings"}]


def test_parse_table_with_header_separator_and_body_rows():
    md = (
        "| Name | Dtype |\n"
        "| --- | --- |\n"
        "| age | int64 |\n"
        "| income | float64 |"
    )
    blocks = parse_markdown_blocks(md)
    assert blocks == [{
        "type": "table",
        "headers": ["Name", "Dtype"],
        "rows": [["age", "int64"], ["income", "float64"]],
    }]


def test_parse_blockquote():
    blocks = parse_markdown_blocks("> Small clean classification dataset.")
    assert blocks == [{"type": "blockquote", "text": "Small clean classification dataset."}]


def test_parse_bullets_grouped_into_one_block():
    md = "- imputed age\n- dropped dupes"
    blocks = parse_markdown_blocks(md)
    assert blocks == [{"type": "bullet", "items": ["imputed age", "dropped dupes"]}]


def test_parse_none_marker_is_italic():
    blocks = parse_markdown_blocks("_None._")
    assert blocks == [{"type": "italic", "text": "None."}]


def test_parse_stub_line_is_italic():
    blocks = parse_markdown_blocks("_Cleaning has not been run for this dataset._")
    assert blocks == [{
        "type": "italic",
        "text": "Cleaning has not been run for this dataset.",
    }]


def test_parse_generated_timestamp_line_is_italic():
    blocks = parse_markdown_blocks("_Generated 2026-09-01T12:00:00+00:00 (UTC)_")
    assert blocks == [{
        "type": "italic",
        "text": "Generated 2026-09-01T12:00:00+00:00 (UTC)",
    }]


def test_parse_plain_paragraph():
    blocks = parse_markdown_blocks("This is plain text.")
    assert blocks == [{"type": "paragraph", "text": "This is plain text."}]


def test_parse_blank_line_is_spacer():
    blocks = parse_markdown_blocks("para one\n\npara two")
    assert blocks == [
        {"type": "paragraph", "text": "para one"},
        {"type": "spacer"},
        {"type": "paragraph", "text": "para two"},
    ]


def test_parse_mixed_document():
    md = (
        "# Title\n\n"
        "## Section\n\n"
        "**Sub**\n"
        "some intro\n\n"
        "| A | B |\n"
        "| --- | --- |\n"
        "| 1 | 2 |\n\n"
        "> a quote\n\n"
        "- one\n"
        "- two\n\n"
        "_None._"
    )
    blocks = parse_markdown_blocks(md)
    types = [b["type"] for b in blocks]
    assert types == [
        "title", "spacer",
        "heading", "spacer",
        "subheading", "paragraph", "spacer",
        "table", "spacer",
        "blockquote", "spacer",
        "bullet", "spacer",
        "italic",
    ]


# ---------------------------------------------------------------------------
# Helper: build a real FinalReport for export tests
# ---------------------------------------------------------------------------
def _build_final_report() -> FinalReport:
    save_artifact("ds-export", "understanding", FAKE_UNDERSTANDING)
    save_artifact("ds-export", "cleaning", FAKE_CLEANING)
    save_artifact("ds-export", "training", FAKE_TRAINING)
    client = FakeAnthropicClient([_response([_tool_use_block("submit_report_prose", PROSE)])])
    return run_report("ds-export", client=client)


def test_render_pdf_returns_nonempty_bytes_with_pdf_magic_header():
    fr = _build_final_report()
    pdf_bytes = render_pdf(fr)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 100
    assert pdf_bytes.startswith(b"%PDF")


def test_render_docx_returns_nonempty_bytes_with_zip_magic_header():
    fr = _build_final_report()
    docx_bytes = render_docx(fr)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 100
    assert docx_bytes.startswith(b"PK")


def test_render_docx_contains_real_facts_from_sidecars():
    fr = _build_final_report()
    docx_bytes = render_docx(fr)
    doc = Document(io.BytesIO(docx_bytes))

    all_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += "\n" + cell.text

    assert "xgboost" in all_text          # best model name, verbatim
    assert "impute_column" in all_text    # cleaning step tool name, verbatim
    assert "0.91" in all_text             # correlation value, verbatim


# ---------------------------------------------------------------------------
# Endpoint tests
# ---------------------------------------------------------------------------
@pytest.fixture
def report_client(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENTDS_DATA_DIR", str(tmp_path / "uploads"))
    from app.main import app
    from fastapi.testclient import TestClient

    tc = TestClient(app)
    csv = b"age,income,target\n25,50000,0\n30,60000,1\n41,72000,1\n22,,0\n"
    ds = tc.post(
        "/datasets/upload",
        files={"file": ("people.csv", io.BytesIO(csv), "text/csv")},
    ).json()["dataset_id"]
    return tc, ds


def test_get_report_pdf_404_before_build(report_client):
    tc, ds = report_client
    resp = tc.get(f"/datasets/{ds}/report/pdf")
    assert resp.status_code == 404


def test_get_report_docx_404_before_build(report_client):
    tc, ds = report_client
    resp = tc.get(f"/datasets/{ds}/report/docx")
    assert resp.status_code == 404


def test_get_report_pdf_after_post_returns_pdf(report_client, monkeypatch):
    tc, ds = report_client
    monkeypatch.setenv("AGENTDS_LLM_PROVIDER", "anthropic")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    def fake_du_run(self, df, ds_id):
        from app.agents.data_understanding import ColumnOverview, DataUnderstandingReport, list_columns
        return DataUnderstandingReport(
            dataset_id=ds_id, n_rows=len(df), n_columns=df.shape[1],
            columns=[ColumnOverview(**c) for c in list_columns(df)],
            problem_type="classification", target_candidate="target",
            reasoning="stub", narrative="stub narrative", key_findings=[],
        )

    def fake_run_report(dataset_id, client=None):
        return run_report(dataset_id, client=FakeAnthropicClient(
            [_response([_tool_use_block("submit_report_prose", PROSE)])]))

    monkeypatch.setattr(
        "app.agents.data_understanding.DataUnderstandingAgent.run", fake_du_run
    )
    monkeypatch.setattr("app.routers.report.run_report", fake_run_report)
    posted = tc.post(f"/datasets/{ds}/report")
    assert posted.status_code == 200

    resp = tc.get(f"/datasets/{ds}/report/pdf")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")


def test_get_report_docx_after_post_returns_docx(report_client, monkeypatch):
    tc, ds = report_client
    monkeypatch.setenv("AGENTDS_LLM_PROVIDER", "anthropic")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    def fake_du_run(self, df, ds_id):
        from app.agents.data_understanding import ColumnOverview, DataUnderstandingReport, list_columns
        return DataUnderstandingReport(
            dataset_id=ds_id, n_rows=len(df), n_columns=df.shape[1],
            columns=[ColumnOverview(**c) for c in list_columns(df)],
            problem_type="classification", target_candidate="target",
            reasoning="stub", narrative="stub narrative", key_findings=[],
        )

    def fake_run_report(dataset_id, client=None):
        return run_report(dataset_id, client=FakeAnthropicClient(
            [_response([_tool_use_block("submit_report_prose", PROSE)])]))

    monkeypatch.setattr(
        "app.agents.data_understanding.DataUnderstandingAgent.run", fake_du_run
    )
    monkeypatch.setattr("app.routers.report.run_report", fake_run_report)
    posted = tc.post(f"/datasets/{ds}/report")
    assert posted.status_code == 200

    resp = tc.get(f"/datasets/{ds}/report/docx")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert resp.content.startswith(b"PK")


def test_export_endpoints_work_without_api_key_once_report_is_built(report_client, monkeypatch):
    tc, ds = report_client
    monkeypatch.setenv("AGENTDS_LLM_PROVIDER", "anthropic")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    def fake_du_run(self, df, ds_id):
        from app.agents.data_understanding import ColumnOverview, DataUnderstandingReport, list_columns
        return DataUnderstandingReport(
            dataset_id=ds_id, n_rows=len(df), n_columns=df.shape[1],
            columns=[ColumnOverview(**c) for c in list_columns(df)],
            problem_type="classification", target_candidate="target",
            reasoning="stub", narrative="stub narrative", key_findings=[],
        )

    def fake_run_report(dataset_id, client=None):
        return run_report(dataset_id, client=FakeAnthropicClient(
            [_response([_tool_use_block("submit_report_prose", PROSE)])]))

    monkeypatch.setattr(
        "app.agents.data_understanding.DataUnderstandingAgent.run", fake_du_run
    )
    monkeypatch.setattr("app.routers.report.run_report", fake_run_report)
    posted = tc.post(f"/datasets/{ds}/report")
    assert posted.status_code == 200

    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)

    pdf_resp = tc.get(f"/datasets/{ds}/report/pdf")
    docx_resp = tc.get(f"/datasets/{ds}/report/docx")
    assert pdf_resp.status_code == 200
    assert pdf_resp.content.startswith(b"%PDF")
    assert docx_resp.status_code == 200
    assert docx_resp.content.startswith(b"PK")
